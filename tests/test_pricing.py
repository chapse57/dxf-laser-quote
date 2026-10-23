"""Pricing checked by hand against the formula in pricing.py."""
import ezdxf
import pytest

from laserquote.geometry import measure
from laserquote.pricing import load_rates, quote, write_default_rates


@pytest.fixture
def rates(tmp_path):
    return load_rates(write_default_rates(tmp_path / "rates.xlsx"))


def plate_100x50():
    doc = ezdxf.new(units=4)
    msp = doc.modelspace()
    pts = [(0, 0), (100, 0), (100, 50), (0, 50)]
    for a, b in zip(pts, pts[1:] + pts[:1]):
        msp.add_line(a, b)
    return measure(doc)  # 300 mm, 1 pierce, 5000 mm2 blank


def test_default_sheet_loads(rates):
    assert "Mild steel" in rates.material_names
    assert rates.thicknesses("Mild steel") == [1.0, 2.0, 3.0, 5.0, 6.0, 10.0]
    assert rates.settings["currency"] == "USD"


def test_price_matches_hand_calculation(rates):
    q = quote(plate_100x50(), rates, "Mild steel", 3.0, quantity=10)
    # Mild steel 3 mm: 4000 mm/min, 1.0 s pierce, 80 per m2
    cut_min = 300 / 4000  # 0.075
    pierce_min = 1 * 1.0 / 60
    machine = (cut_min + pierce_min) / 60 * 90
    material = 5000 / 1e6 * 1.15 * 80
    part = machine + material
    subtotal = part * 10 + 15
    total = subtotal * 1.20
    assert q.cut_time_min == pytest.approx(cut_min)
    assert q.machine_cost == pytest.approx(machine)
    assert q.material_cost == pytest.approx(material)
    assert q.subtotal == pytest.approx(subtotal)
    assert q.total == pytest.approx(round(total, 2))
    assert not q.min_order_applied


def test_minimum_order_kicks_in_for_a_tiny_part(rates):
    doc = ezdxf.new(units=4)
    doc.modelspace().add_circle((0, 0), 2)  # 12.6 mm of cut
    q = quote(measure(doc), rates, "Mild steel", 1.0, quantity=1)
    assert q.min_order_applied
    assert q.total == 25.0


def test_unknown_material_is_an_error(rates):
    with pytest.raises(KeyError):
        quote(plate_100x50(), rates, "Titanium", 3.0)


def test_edited_sheet_changes_the_price(tmp_path, rates):
    import openpyxl

    path = write_default_rates(tmp_path / "r.xlsx")
    wb = openpyxl.load_workbook(path)
    st = wb["Settings"]
    for row in st.iter_rows(min_row=2):
        if row[0].value == "margin_pct":
            row[1].value = 0
        if row[0].value == "setup_fee":
            row[1].value = 0
    wb.save(path)
    edited = load_rates(path)
    q = quote(plate_100x50(), edited, "Mild steel", 3.0, quantity=100)
    assert not q.min_order_applied
    assert q.total == pytest.approx(round(q.part_cost * 100, 2))


def test_bad_number_in_sheet_gives_a_clear_message(tmp_path):
    import openpyxl

    path = write_default_rates(tmp_path / "r.xlsx")
    wb = openpyxl.load_workbook(path)
    wb["Materials"]["C2"] = "fast"
    wb.save(path)
    with pytest.raises(ValueError, match="row 2.*cut_speed_mm_per_min"):
        load_rates(path)
