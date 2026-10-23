"""Turn a measurement into a price using the shop's own Excel rate sheet.

rates.xlsx has two sheets the shop edits itself, no code involved:

  Materials  - one row per material + thickness:
               material, thickness_mm, cut_speed_mm_per_min, pierce_sec,
               price_per_m2 (material cost per square metre of sheet)
  Settings   - name / value pairs:
               machine_rate_per_hour, setup_fee, min_order, margin_pct,
               material_waste_pct, currency

How the price is built (every step is shown in the breakdown):

  cut time       = cut length / cut speed  +  pierces x pierce time
  machine cost   = cut time x machine rate
  material cost  = blank area x (1 + waste) x price per m2
  per part       = machine + material
  order          = (per part x qty + setup fee) x (1 + margin)
  final          = max(order, min_order)
"""
from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Union

import openpyxl

from .geometry import Measurement

MATERIAL_COLUMNS = ["material", "thickness_mm", "cut_speed_mm_per_min", "pierce_sec", "price_per_m2"]
SETTING_DEFAULTS = {
    "machine_rate_per_hour": 90.0,
    "setup_fee": 15.0,
    "min_order": 25.0,
    "margin_pct": 20.0,
    "material_waste_pct": 15.0,
    "currency": "USD",
}


@dataclass(frozen=True)
class MaterialRate:
    material: str
    thickness_mm: float
    cut_speed_mm_per_min: float
    pierce_sec: float
    price_per_m2: float

    @property
    def label(self) -> str:
        t = f"{self.thickness_mm:g}"
        return f"{self.material} {t} mm"


@dataclass
class RateSheet:
    materials: list[MaterialRate]
    settings: dict
    source: str

    @property
    def material_names(self) -> list[str]:
        seen: list[str] = []
        for m in self.materials:
            if m.material not in seen:
                seen.append(m.material)
        return seen

    def thicknesses(self, material: str) -> list[float]:
        return sorted(m.thickness_mm for m in self.materials if m.material == material)

    def get(self, material: str, thickness_mm: float) -> MaterialRate:
        for m in self.materials:
            if m.material == material and abs(m.thickness_mm - thickness_mm) < 1e-9:
                return m
        raise KeyError(f"No rate for {material} {thickness_mm} mm in {self.source}")


@dataclass
class Quote:
    rate: MaterialRate
    quantity: int
    cut_length_mm: float
    pierces: int
    blank_area_mm2: float
    cut_time_min: float  # per part, cutting only
    pierce_time_min: float  # per part
    machine_cost: float  # per part
    material_cost: float  # per part
    part_cost: float  # per part, before setup and margin
    setup_fee: float
    subtotal: float  # all parts + setup, before margin
    margin: float
    min_order_applied: bool
    total: float
    currency: str

    @property
    def time_per_part_min(self) -> float:
        return self.cut_time_min + self.pierce_time_min

    @property
    def unit_price(self) -> float:
        return self.total / self.quantity if self.quantity else 0.0

    def rows(self) -> list[tuple[str, str]]:
        """Breakdown lines for display, in order."""
        c = self.currency
        return [
            ("Material", f"{self.rate.label}"),
            ("Cut length", f"{self.cut_length_mm / 1000:.3f} m at {self.rate.cut_speed_mm_per_min:g} mm/min"),
            ("Cutting time", f"{self.cut_time_min:.2f} min"),
            ("Pierces", f"{self.pierces} x {self.rate.pierce_sec:g} s = {self.pierce_time_min:.2f} min"),
            ("Machine cost per part", f"{self.machine_cost:.2f} {c}"),
            ("Material per part", f"{self.blank_area_mm2 / 1e6:.4f} m² blank (+waste) = {self.material_cost:.2f} {c}"),
            ("Cost per part", f"{self.part_cost:.2f} {c}"),
            ("Quantity", f"{self.quantity}"),
            ("Setup fee", f"{self.setup_fee:.2f} {c}"),
            ("Subtotal", f"{self.subtotal:.2f} {c}"),
            ("Margin", f"{self.margin:.2f} {c}"),
            ("Total", f"{self.total:.2f} {c}" + ("  (minimum order)" if self.min_order_applied else "")),
            ("Unit price", f"{self.unit_price:.2f} {c}"),
        ]


# ------------------------------------------------------------ rate sheet

def _num(v, name: str, row: int) -> float:
    try:
        return float(v)
    except (TypeError, ValueError):
        raise ValueError(f"Materials sheet row {row}: '{name}' must be a number, got {v!r}") from None


def load_rates(source: Union[str, Path, bytes]) -> RateSheet:
    if isinstance(source, bytes):
        import io
        wb = openpyxl.load_workbook(io.BytesIO(source), data_only=True)
        name = "uploaded rates.xlsx"
    else:
        wb = openpyxl.load_workbook(source, data_only=True)
        name = str(source)

    if "Materials" not in wb.sheetnames:
        raise ValueError("rates.xlsx needs a sheet called 'Materials'")
    ws = wb["Materials"]
    header = [str(c.value).strip().lower() if c.value is not None else "" for c in ws[1]]
    missing = [c for c in MATERIAL_COLUMNS if c not in header]
    if missing:
        raise ValueError(f"Materials sheet is missing columns: {', '.join(missing)}")
    col = {c: header.index(c) for c in MATERIAL_COLUMNS}

    materials: list[MaterialRate] = []
    for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if row is None or all(v in (None, "") for v in row):
            continue
        mat = str(row[col["material"]]).strip()
        if not mat or mat.lower() == "none":
            continue
        speed = _num(row[col["cut_speed_mm_per_min"]], "cut_speed_mm_per_min", i)
        if speed <= 0:
            raise ValueError(f"Materials sheet row {i}: cut speed must be above 0")
        materials.append(
            MaterialRate(
                material=mat,
                thickness_mm=_num(row[col["thickness_mm"]], "thickness_mm", i),
                cut_speed_mm_per_min=speed,
                pierce_sec=_num(row[col["pierce_sec"]], "pierce_sec", i),
                price_per_m2=_num(row[col["price_per_m2"]], "price_per_m2", i),
            )
        )
    if not materials:
        raise ValueError("Materials sheet has no rows")

    settings = dict(SETTING_DEFAULTS)
    if "Settings" in wb.sheetnames:
        for row in wb["Settings"].iter_rows(min_row=1, values_only=True):
            if not row or row[0] is None:
                continue
            key = str(row[0]).strip()
            if key in SETTING_DEFAULTS and len(row) > 1 and row[1] is not None:
                if key == "currency":
                    settings[key] = str(row[1]).strip()
                else:
                    settings[key] = float(row[1])
    return RateSheet(materials, settings, name)


# ------------------------------------------------------------- quoting

def quote(m: Measurement, rates: RateSheet, material: str, thickness_mm: float, quantity: int = 1) -> Quote:
    if quantity < 1:
        raise ValueError("quantity must be at least 1")
    rate = rates.get(material, thickness_mm)
    s = rates.settings

    cut_time_min = m.cut_length_mm / rate.cut_speed_mm_per_min
    pierce_time_min = m.pierces * rate.pierce_sec / 60
    machine_cost = (cut_time_min + pierce_time_min) / 60 * s["machine_rate_per_hour"]
    material_cost = m.blank_area_mm2 / 1e6 * (1 + s["material_waste_pct"] / 100) * rate.price_per_m2
    part_cost = machine_cost + material_cost

    subtotal = part_cost * quantity + s["setup_fee"]
    margin = subtotal * s["margin_pct"] / 100
    total = subtotal + margin
    min_applied = total < s["min_order"]
    if min_applied:
        total = s["min_order"]

    return Quote(
        rate=rate,
        quantity=quantity,
        cut_length_mm=m.cut_length_mm,
        pierces=m.pierces,
        blank_area_mm2=m.blank_area_mm2,
        cut_time_min=cut_time_min,
        pierce_time_min=pierce_time_min,
        machine_cost=machine_cost,
        material_cost=material_cost,
        part_cost=part_cost,
        setup_fee=s["setup_fee"],
        subtotal=subtotal,
        margin=margin,
        min_order_applied=min_applied,
        total=round(total, 2),
        currency=s["currency"],
    )


# ------------------------------------------------------ default sheet

DEFAULT_MATERIALS = [
    # material, thickness, speed mm/min, pierce s, price per m2 (placeholder numbers)
    ("Mild steel", 1.0, 9000, 0.5, 30),
    ("Mild steel", 2.0, 6000, 0.8, 55),
    ("Mild steel", 3.0, 4000, 1.0, 80),
    ("Mild steel", 5.0, 2200, 1.5, 130),
    ("Mild steel", 6.0, 1800, 2.0, 155),
    ("Mild steel", 10.0, 900, 3.5, 260),
    ("Stainless 304", 1.0, 7000, 0.6, 90),
    ("Stainless 304", 2.0, 4000, 1.0, 170),
    ("Stainless 304", 3.0, 2500, 1.5, 250),
    ("Stainless 304", 5.0, 1300, 2.5, 410),
    ("Aluminium 5052", 1.0, 8000, 0.5, 60),
    ("Aluminium 5052", 2.0, 5000, 0.8, 110),
    ("Aluminium 5052", 3.0, 3200, 1.2, 160),
    ("Aluminium 5052", 5.0, 1600, 2.0, 265),
]


def write_default_rates(path: Union[str, Path]) -> Path:
    """Create a starter rates.xlsx the shop can edit. Numbers are placeholders."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Materials"
    ws.append(MATERIAL_COLUMNS)
    for row in DEFAULT_MATERIALS:
        ws.append(list(row))
    for c, w in zip("ABCDE", (18, 14, 22, 12, 14)):
        ws.column_dimensions[c].width = w

    st = wb.create_sheet("Settings")
    st.append(["name", "value", "what it means"])
    notes = {
        "machine_rate_per_hour": "what one hour of laser time costs you, incl. gas and operator",
        "setup_fee": "flat fee per order for loading the sheet and the program",
        "min_order": "smallest total you will invoice",
        "margin_pct": "profit added on top of cost, in percent",
        "material_waste_pct": "extra sheet used up by nesting gaps and offcuts, in percent",
        "currency": "shown next to every price",
    }
    for k, v in SETTING_DEFAULTS.items():
        st.append([k, v, notes[k]])
    st.column_dimensions["A"].width = 24
    st.column_dimensions["C"].width = 64

    path = Path(path)
    wb.save(path)
    return path
